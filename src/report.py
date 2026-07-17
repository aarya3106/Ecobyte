"""
Report Module
Generates PDF, DOCX, and CSV reports summarizing the climate risk assessment and recommendations.
"""

import os
import sys
import io
import pandas as pd
from fpdf import FPDF
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_pdf_report(city_data: dict, recommendations: dict, ward_name: str = "N/A") -> bytes:
    """
    Generates a beautifully formatted PDF report for the active city.
    
    Args:
        city_data (dict): Current twin snapshot database parameters.
        recommendations (dict): Grouped policy recommendations.
        ward_name (str): Selected target ward name.
        
    Returns:
        bytes: Rendered PDF file data.
    """
    # Locate fonts directory relative to the project root
    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(current_dir)
    fonts_dir = os.path.join(project_root, "fonts")
    
    regular_path = os.path.join(fonts_dir, "NotoSans-Regular.ttf")
    bold_path = os.path.join(fonts_dir, "NotoSans-Bold.ttf")
    italic_path = os.path.join(fonts_dir, "NotoSans-Italic.ttf")
    bold_italic_path = os.path.join(fonts_dir, "NotoSans-BoldItalic.ttf")
    
    pdf = FPDF()
    
    # Check if font files exist and register them
    has_unicode_font = False
    font_family = "helvetica" # Fallback core font
    
    if (os.path.exists(regular_path) and 
        os.path.exists(bold_path) and 
        os.path.exists(italic_path) and 
        os.path.exists(bold_italic_path)):
        try:
            pdf.add_font("NotoSans", style="", fname=regular_path)
            pdf.add_font("NotoSans", style="B", fname=bold_path)
            pdf.add_font("NotoSans", style="I", fname=italic_path)
            pdf.add_font("NotoSans", style="BI", fname=bold_italic_path)
            font_family = "NotoSans"
            has_unicode_font = True
        except Exception as e:
            print(f"[Report PDF Error] Failed to register NotoSans font: {e}", file=sys.stderr)
            
    pdf.add_page()
    
    def clean_text(text: str) -> str:
        # Safety net fallback: if Unicode font is not loaded, replace Indian Rupee symbol to prevent crash
        if not has_unicode_font:
            return text.replace("₹", "Rs. ")
        return text

    # Define Document Typography
    pdf.set_font(font_family, "B", 20)
    
    # Header Accent Bar
    pdf.set_fill_color(0, 173, 181) # Teal accent
    pdf.rect(10, 10, 190, 8, 'F')
    pdf.ln(12)
    
    # Document Title
    pdf.set_text_color(30, 30, 30)
    pdf.cell(0, 12, clean_text(f"Climate Risk Report - {ward_name}, {city_data['city']}"), ln=True, align='C')
    
    # Metadata Subtitle
    pdf.set_font(font_family, "I", 9)
    pdf.set_text_color(120, 120, 120)
    last_updated_str = city_data.get("last_updated", "N/A")[:10]
    pdf.cell(0, 8, clean_text(f"Assessment Date: {last_updated_str} | Target Zone: {ward_name}"), ln=True, align='C')
    
    # Methodology disclosure note
    pdf.set_text_color(230, 126, 34)
    pdf.set_font(font_family, "I", 8.5)
    pdf.cell(0, 5, clean_text("Ward-level figures are illustrative estimates for demonstration purposes;"), ln=True, align='C')
    pdf.cell(0, 5, clean_text("underlying models are trained on city-level data."), ln=True, align='C')
    pdf.ln(8)
    
    # Section 1: Executive Risk Profile
    pdf.set_font(font_family, "B", 13)
    pdf.set_text_color(30, 30, 30)
    pdf.cell(0, 8, clean_text("1. Executive Risk Profile"), ln=True)
    
    # Colored underline for section header
    pdf.set_draw_color(0, 173, 181)
    pdf.set_line_width(0.5)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    # Risk summary paragraph
    pdf.set_font(font_family, "", 10)
    pdf.set_text_color(60, 60, 60)
    summary_text = (
        f"This report outlines the forecasted climate risk indices for the city of {city_data['city']} "
        f"derived from active digital twin parameters: Average Temperature is {city_data['avg_temp']:.1f}°C, "
        f"Rainfall is {city_data['precipitation']:.1f}mm, Humidity is {city_data['humidity']:.1f}%, and "
        f"Population is {city_data['population']:,}. These features have been evaluated by the "
        f"IMD-calibrated EcoByte engine to forecast risk levels across three key dimensions."
    )
    pdf.multi_cell(0, 5, clean_text(summary_text))
    pdf.ln(6)
    
    # Risk Scores Grid Table
    pdf.set_font(font_family, "B", 10)
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(30, 30, 30) # Dark header
    pdf.cell(60, 8, clean_text("Risk Indicator"), border=1, align='C', fill=True)
    pdf.cell(60, 8, clean_text("Decision Engine Score (0-100)"), border=1, align='C', fill=True)
    pdf.cell(60, 8, clean_text("Risk Category Band"), border=1, align='C', fill=True)
    pdf.ln()
    
    pdf.set_text_color(60, 60, 60)
    pdf.set_font(font_family, "", 9)
    
    hazards = [
        ("heat_risk", "heat_risk_score", "Heat Wave Risk"),
        ("flood_risk", "flood_risk_score", "Urban Flood Risk"),
        ("water_stress", "water_stress_score", "Water Stress Index")
    ]
    
    for key, score_key, name in hazards:
        score = city_data[score_key]
        level = recommendations[key]["level"]
        
        # Color coding for Risk Level
        if level == "Low":
            pdf.set_text_color(46, 204, 113) # Green
        elif level == "Moderate":
            pdf.set_text_color(241, 196, 15) # Yellow
        elif level == "High":
            pdf.set_text_color(230, 126, 34) # Orange
        else:
            pdf.set_text_color(231, 76, 60) # Red
            
        pdf.cell(60, 8, clean_text(name), border=1, align='C')
        pdf.cell(60, 8, f"{score:.2f}", border=1, align='C')
        pdf.cell(60, 8, clean_text(level), border=1, align='C')
        pdf.ln()
        
    pdf.ln(10)
    
    # Section 2: Adaptation Actions
    pdf.set_text_color(30, 30, 30)
    pdf.set_font(font_family, "B", 13)
    pdf.cell(0, 8, clean_text("2. Targeted Policy Adaptation Checklists"), ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    adaptation_groups = [
        ("heat_risk", "Heat Wave Mitigation Responses"),
        ("flood_risk", "Stormwater & Flood Management"),
        ("water_stress", "Water Conservation & Supply Assurance")
    ]
    
    for key, title in adaptation_groups:
        pdf.set_font(font_family, "B", 11)
        pdf.set_text_color(0, 173, 181)
        pdf.cell(0, 6, clean_text(title), ln=True)
        pdf.ln(1)
        
        pdf.set_font(font_family, "", 9.5)
        pdf.set_text_color(60, 60, 60)
        
        actions = recommendations[key]["actions"]
        if not actions:
            pdf.cell(0, 5, clean_text("- No urgent intervention required under the current low-risk threshold."), ln=True)
            pdf.ln(2)
        else:
            for action in actions:
                # Combine Action title and description with cost and timeline metadata
                meta_tag = f" [{action['impact']} | Cost: {action.get('cost','N/A')} | Time: {action.get('timeline','N/A')}]: "
                pdf.set_font(font_family, "B", 9)
                pdf.write(5, clean_text(f"-  {action['action'].upper()}"))
                pdf.set_font(font_family, "BI", 8.5)
                pdf.set_text_color(100, 110, 120)
                pdf.write(5, clean_text(meta_tag))
                pdf.set_text_color(60, 60, 60)
                pdf.set_font(font_family, "", 9)
                pdf.write(5, clean_text(f"{action['rationale']}\n"))
            pdf.ln(3)
            
    # Add Data Provenance Footer
    pdf.ln(5)
    pdf.set_font(font_family, "I", 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 4, clean_text("Data Sources & Provenance Note:"), ln=True)
    pdf.cell(0, 4, clean_text("- Temperature & Humidity: Open-Meteo real-time station records & historical archive."), ln=True)
    pdf.cell(0, 4, clean_text("- Precipitation: Meteostat station telemetry & Open-Meteo monsoon tracking."), ln=True)
    pdf.cell(0, 4, clean_text("- Population: Census of India 2011 official municipal ward boundaries."), ln=True)
    pdf.cell(0, 4, clean_text("- Water Supply Limits: MCGM / PMC / NMC Annual Hydraulic Reports (2022-2023)."), ln=True)
    
    # Output the PDF as bytes
    return bytes(pdf.output())

def generate_word_report(city_data: dict, recommendations: dict, ward_name: str = "N/A") -> bytes:
    """
    Generates a cleanly formatted Microsoft Word (DOCX) report for the active city.
    
    Args:
        city_data (dict): Current twin snapshot database parameters.
        recommendations (dict): Grouped policy recommendations.
        ward_name (str): Selected target ward name.
        
    Returns:
        bytes: Rendered DOCX file data.
    """
    doc = docx.Document()
    
    # Title Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(f"Climate Risk Report — {ward_name}, {city_data['city']}")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    
    # Metadata Subtitle
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    last_updated_str = city_data.get("last_updated", "N/A")[:10]
    run_meta = p_meta.add_run(f"Generated on: {last_updated_str} | Target Zone: {ward_name}")
    run_meta.font.name = 'Arial'
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    
    # Methodology disclosure note (orange colored warning text)
    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_disc = p_disc.add_run(
        "Ward-level figures are illustrative estimates for demonstration purposes; "
        "underlying models are trained on city-level data."
    )
    run_disc.font.name = 'Arial'
    run_disc.font.size = Pt(9.5)
    run_disc.font.italic = True
    run_disc.font.color.rgb = docx.shared.RGBColor(230, 126, 34)
    
    doc.add_paragraph() # Spacing
    
    # Section 1: Executive Profile
    p_h1 = doc.add_paragraph()
    run_h1 = p_h1.add_run("1. Executive Risk Profile")
    run_h1.font.name = 'Arial'
    run_h1.font.size = Pt(14)
    run_h1.font.bold = True
    
    # Summary Paragraph
    p_sum = doc.add_paragraph()
    run_sum = p_sum.add_run(
        f"This assessment represents the localized climate risk indices for {city_data['city']} "
        f"compiled from raw meteorological twin readings: Average Temperature ({city_data['avg_temp']:.1f}°C), "
        f"Monsoon Rainfall ({city_data['precipitation']:.1f}mm), Relative Humidity ({city_data['humidity']:.1f}%), "
        f"and Population density markers ({city_data['population']:,}). "
        f"The IMD-calibrated EcoByte engine uses these indicators to calculate the following safety indexes."
    )
    run_sum.font.name = 'Arial'
    run_sum.font.size = Pt(11)
    
    # Table of risk scores
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Light Shading Accent 1'
    
    # Format Headers
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Risk Dimension'
    hdr_cells[1].text = 'Engine Score (0-100)'
    hdr_cells[2].text = 'Risk Level Band'
    
    # Fill Data
    hazards = [
        ("heat_risk", "heat_risk_score", "Heat Wave Risk"),
        ("flood_risk", "flood_risk_score", "Urban Flood Risk"),
        ("water_stress", "water_stress_score", "Water Stress Index")
    ]
    for idx, (key, score_key, name) in enumerate(hazards, 1):
        row_cells = table.rows[idx].cells
        row_cells[0].text = name
        row_cells[1].text = f"{city_data[score_key]:.2f}"
        row_cells[2].text = recommendations[key]["level"]
        
    doc.add_paragraph() # Spacing
    
    # Section 2: Adaptation Actions
    p_h2 = doc.add_paragraph()
    run_h2 = p_h2.add_run("2. Targeted Policy Adaptation Checklists")
    run_h2.font.name = 'Arial'
    run_h2.font.size = Pt(14)
    run_h2.font.bold = True
    
    adaptation_groups = [
        ("heat_risk", "Heat Wave Mitigation Responses"),
        ("flood_risk", "Stormwater & Flood Management"),
        ("water_stress", "Water Conservation & Supply Assurance")
    ]
    
    for key, title in adaptation_groups:
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(title)
        run_sub.font.name = 'Arial'
        run_sub.font.size = Pt(12)
        run_sub.font.bold = True
        
        actions = recommendations[key]["actions"]
        if not actions:
            p_action = doc.add_paragraph(style='List Bullet')
            p_action.add_run("No urgent intervention required under current low-risk threshold.")
        else:
            for action in actions:
                p_action = doc.add_paragraph(style='List Bullet')
                r_act = p_action.add_run(f"{action['action']} ")
                r_act.font.bold = True
                r_act.font.name = 'Arial'
                
                meta_str = f"[{action['impact']} | Cost: {action.get('cost','N/A')} | Time: {action.get('timeline','N/A')}]: "
                r_tag = p_action.add_run(meta_str)
                r_tag.font.italic = True
                r_tag.font.name = 'Arial'
                
                r_desc = p_action.add_run(action['rationale'])
                r_desc.font.name = 'Arial'
                
    # Add Data Provenance Footer in Word
    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    r_foot_title = p_foot.add_run("Data Sources & Provenance Footnotes:\n")
    r_foot_title.font.bold = True
    r_foot_title.font.size = Pt(9.5)
    r_foot_title.font.name = 'Arial'
    
    footnotes = (
        "- Temperature & Humidity: Open-Meteo real-time station records & historical archive.\n"
        "- Precipitation: Meteostat station telemetry & Open-Meteo monsoon tracking.\n"
        "- Population: Census of India 2011 official municipal ward boundaries.\n"
        "- Water Supply Limits: MCGM / PMC / NMC Annual Hydraulic Reports (2022-2023)."
    )
    r_foot_body = p_foot.add_run(footnotes)
    r_foot_body.font.size = Pt(8.5)
    r_foot_body.font.italic = True
    r_foot_body.font.name = 'Arial'
    
    # Save the doc to a binary IO stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()

def generate_csv_export(city_data: dict, ward_name: str = "N/A") -> bytes:
    """
    Converts a single city's snapshot dictionary into a downloadable CSV file.
    
    Args:
        city_data (dict): Current twin snapshot database parameters.
        ward_name (str): Selected target ward name.
        
    Returns:
        bytes: Encoded CSV text data.
    """
    data_copy = city_data.copy()
    data_copy["ward"] = ward_name
    df = pd.DataFrame([data_copy])
    file_stream = io.StringIO()
    df.to_csv(file_stream, index=False)
    return file_stream.getvalue().encode('utf-8')
